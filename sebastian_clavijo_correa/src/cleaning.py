"""Limpia los datos de EA1 y guarda las salidas y los reportes de EA2."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
from pathlib import Path
import sys
from datetime import datetime, timezone
import unicodedata

import duckdb
import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
REQUIRED_COLUMNS = [
    "city_id",
    "time_id",
    "observation_time",
    "city",
    "latitude",
    "longitude",
    "temperature_c",
    "humidity_pct",
    "precipitation_mm",
    "wind_speed_kmh",
    "energy_kwh",
    "synthetic_energy",
]
MEASURES = ["temperature_c", "humidity_pct", "precipitation_mm", "wind_speed_kmh", "energy_kwh"]
IMPUTABLE = ["temperature_c", "humidity_pct", "wind_speed_kmh", "energy_kwh"]
BOUNDS = {
    "latitude": (-90, 90),
    "longitude": (-180, 180),
    "temperature_c": (-90, 60),
    "humidity_pct": (0, 100),
    "precipitation_mm": (0, None),
    "wind_speed_kmh": (0, None),
    "energy_kwh": (0, None),
}
CITY_NAMES = {
    name.casefold(): name for name in ["Bogota", "Medellin", "Cali", "Barranquilla", "Bucaramanga"]
}
TZ = "America/Bogota"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize_city(value):
    if pd.isna(value):
        return pd.NA
    normalized = "".join(
        c for c in unicodedata.normalize("NFKD", str(value).strip()) if not unicodedata.combining(c)
    ).casefold()
    return CITY_NAMES.get(normalized, pd.NA)


def profile(frame: pd.DataFrame) -> dict:
    stats = {}
    for col in MEASURES:
        valid = pd.to_numeric(frame[col], errors="coerce").replace([np.inf, -np.inf], np.nan)
        summary = valid.describe(percentiles=[0.25, 0.5, 0.75])
        stats[col] = {str(k): None if pd.isna(v) else float(v) for k, v in summary.items()}
    return {
        "rows": len(frame),
        "columns": len(frame.columns),
        "nulls": {col: int(frame[col].isna().sum()) for col in frame.columns},
        "dtypes": {col: str(dtype) for col, dtype in frame.dtypes.items()},
        "exact_duplicate_rows": int(frame.duplicated(subset=REQUIRED_COLUMNS).sum()),
        "duplicate_key_rows": int(frame.duplicated(["city", "observation_time"], keep=False).sum()),
        "statistics": stats,
    }


def extract_database(path: Path) -> tuple[pd.DataFrame, dict]:
    """Consulta los hechos y sus dimensiones sin modificar la base de origen."""
    if not path.is_file():
        raise FileNotFoundError(f"No existe la base de EA1: {path}")
    source_hash = sha256(path)
    with duckdb.connect(str(path), read_only=True) as con:
        tables = {row[0] for row in con.execute("SHOW TABLES").fetchall()}
        expected = {"dim_city", "dim_time", "fact_weather_energy", "staging_weather"}
        if not expected <= tables:
            raise ValueError(f"Faltan tablas EA1: {sorted(expected-tables)}")
        # Reviso los ID antes del JOIN para evitar que un duplicado multiplique las filas.
        for table, key in [("dim_city", "city_id"), ("dim_time", "time_id")]:
            total, distinct, nonnull = con.execute(
                f"SELECT count(*),count(distinct {key}),count({key}) FROM {table}"
            ).fetchone()
            if total != distinct or total != nonnull:
                raise ValueError(f"Integridad dimensional inválida en {table}.{key}")
        orphan_count = con.execute("""SELECT count(*) FROM fact_weather_energy f
            LEFT JOIN dim_city c ON c.city_id=f.city_id
            LEFT JOIN dim_time t ON t.time_id=f.time_id
            WHERE c.city_id IS NULL OR t.time_id IS NULL""").fetchone()[0]
        if orphan_count:
            raise ValueError(
                f"{orphan_count} hechos sin dimensión; no se permite perder filas en el JOIN"
            )
        table_counts = {
            table: con.execute(f"SELECT count(*) FROM {table}").fetchone()[0]
            for table in sorted(expected)
        }
        raw = con.execute(
            """SELECT f.city_id,f.time_id,t.observation_time,c.city,c.latitude,c.longitude,
            f.temperature_c,f.humidity_pct,f.precipitation_mm,f.wind_speed_kmh,f.energy_kwh,f.synthetic_energy
            FROM fact_weather_energy f JOIN dim_city c ON c.city_id=f.city_id
            JOIN dim_time t ON t.time_id=f.time_id ORDER BY c.city,t.observation_time"""
        ).df()
        if len(raw) != table_counts["fact_weather_energy"]:
            raise ValueError("El JOIN cambió el número de hechos")
    if sha256(path) != source_hash:
        raise RuntimeError("Cambió la base durante su lectura")
    return raw, {
        "sha256": source_hash,
        "bytes": path.stat().st_size,
        "tables": table_counts,
        "engine": "DuckDB",
        "read_only": True,
        "event_timezone": TZ,
        "weather_provider": "Open-Meteo Historical Weather API; EA1 no fija models",
        "energy_origin": "Sintética en EA1 con semilla 42, no mediciones de consumo",
        "coordinate_semantics": "Coordenadas solicitadas de ciudad, no de la celda meteorológica",
    }


def clean_frame(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Devuelve los datos limpios, las filas rechazadas y las métricas del proceso.

    Trabaja sobre una copia para conservar la entrada y los valores de los rechazos.
    """
    missing = sorted(set(REQUIRED_COLUMNS) - set(raw.columns))
    if missing:
        raise ValueError(f"Columnas obligatorias ausentes: {missing}")
    original = raw.loc[:, REQUIRED_COLUMNS].copy(deep=True).reset_index(drop=True)
    data = original.copy(deep=True)
    metrics = {"before": profile(data), "operations": {}}
    ops = metrics["operations"]
    data["source_row"] = np.arange(1, len(data) + 1)
    raw_duplicates = data.duplicated(subset=REQUIRED_COLUMNS)
    duplicate_rows = data.loc[raw_duplicates, "source_row"].tolist()
    data = data.loc[~raw_duplicates].copy()
    rejected_parts = []

    def reject(mask, reason):
        nonlocal data
        mask = pd.Series(mask, index=data.index).fillna(True)
        source_rows = data.loc[mask, "source_row"]
        removed = original.iloc[(source_rows - 1).to_numpy()].copy()
        if len(removed):
            removed["source_row"] = source_rows.to_numpy()
            removed["rejection_reason"] = reason
            rejected_parts.append(removed)
        data = data.loc[~mask].copy()

    city_before = data["city"].astype("string")
    data["city"] = data["city"].map(normalize_city).astype("string")
    ops["city_values_normalized"] = int((city_before.fillna("") != data["city"].fillna("")).sum())
    original_time_nonnull = data["observation_time"].notna()
    parsed = pd.to_datetime(data["observation_time"], errors="coerce", format="mixed")
    if not isinstance(parsed.dtype, pd.DatetimeTZDtype):
        parsed = parsed.dt.tz_localize(TZ, ambiguous="NaT", nonexistent="NaT")
    else:
        parsed = parsed.dt.tz_convert(TZ)
    data["observation_time"] = parsed
    ops["invalid_datetimes"] = int((original_time_nonnull & parsed.isna()).sum())

    # Si convierto las humedades 150 y 200 a NaN, se pierde su diferencia.
    # Guardo los valores de origen para detectar que son dos registros en conflicto.
    def metric_identity(value):
        if pd.isna(value):
            return ("missing", "")
        try:
            numeric = float(value)
            if np.isfinite(numeric):
                return ("number", numeric)
        except (ValueError, TypeError):
            pass
        return ("invalid", str(value).strip())

    data["_metric_signature"] = data[MEASURES].apply(
        lambda row: tuple(metric_identity(value) for value in row), axis=1
    )
    ops["numeric_parse_failures"] = {}
    ops["out_of_domain_values"] = {}
    for col in ["city_id", "time_id", "latitude", "longitude"] + MEASURES:
        existed = data[col].notna()
        converted = pd.to_numeric(data[col], errors="coerce").replace([np.inf, -np.inf], np.nan)
        ops["numeric_parse_failures"][col] = int((existed & converted.isna()).sum())
        if col in ["city_id", "time_id"]:
            converted = converted.where((converted > 0) & (converted % 1 == 0))
            data[col] = converted.astype("Int64")
        else:
            low, high = BOUNDS[col]
            invalid = converted.lt(low)
            if high is not None:
                invalid = invalid | converted.gt(high)
            ops["out_of_domain_values"][col] = int(invalid.sum())
            data[col] = converted.mask(invalid).astype("float64")
    boolean_map = {"true": True, "false": False, "1": True, "0": False, "1.0": True, "0.0": False}
    data["synthetic_energy"] = (
        data["synthetic_energy"]
        .astype("string")
        .str.strip()
        .str.casefold()
        .map(boolean_map)
        .astype("boolean")
    )
    essential = [
        "city_id",
        "time_id",
        "observation_time",
        "city",
        "latitude",
        "longitude",
        "synthetic_energy",
    ]
    reject(data[essential].isna().any(axis=1), "INVALID_IDENTITY_COORDINATE_OR_PROVENANCE")
    normalized_duplicates = data.duplicated(subset=REQUIRED_COLUMNS + ["_metric_signature"])
    duplicate_rows.extend(data.loc[normalized_duplicates, "source_row"].tolist())
    data = data.loc[~normalized_duplicates].copy()
    ops["exact_duplicates_removed"] = len(duplicate_rows)
    ops["duplicate_source_rows"] = sorted(duplicate_rows)
    conflicts = data.duplicated(["city", "observation_time"], keep=False)
    ops["conflicting_key_rows"] = int(conflicts.sum())
    reject(conflicts, "CONFLICTING_CITY_TIME_KEY")
    data = data.drop(columns="_metric_signature")
    for name in ["year", "month", "day", "hour"]:
        data[name] = getattr(data["observation_time"].dt, name).astype("int64")
    # Una lluvia faltante no equivale a cero: no permite saber si estuvo seco.
    reject(data["precipitation_mm"].isna(), "MISSING_OR_INVALID_PRECIPITATION")
    data["imputed_fields"] = ""
    ops["imputed_cells"] = {}
    for col in IMPUTABLE:
        median = data.groupby(["city", "year", "month"], observed=True)[col].transform("median")
        mask = data[col].isna() & median.notna()
        data.loc[mask, col] = median.loc[mask]
        data.loc[mask, "imputed_fields"] = data.loc[mask, "imputed_fields"] + col + "|"
        ops["imputed_cells"][col] = int(mask.sum())
    reject(data[MEASURES].isna().any(axis=1), "MISSING_METRIC_WITHOUT_GROUP_DONOR")
    data["imputed_fields"] = data["imputed_fields"].str.rstrip("|")
    data["quality_status"] = np.where(data["imputed_fields"].eq(""), "VALID", "IMPUTED")
    data["is_outlier"] = False
    data["outlier_fields"] = ""
    metrics["outliers"] = {
        "method": "Tukey 1.5 IQR por ciudad; IQR cero no se clasifica",
        "policy": "Señalar y conservar valores físicamente válidos",
        "bounds": [],
        "counts": {},
    }
    for col in MEASURES:
        flag = pd.Series(False, index=data.index)
        for city, group in data.groupby("city", sort=True, observed=True):
            q1, q3 = group[col].quantile([0.25, 0.75])
            iqr = q3 - q1
            lower, upper = q1 - 1.5 * iqr, q3 + 1.5 * iqr
            metrics["outliers"]["bounds"].append(
                {
                    "city": str(city),
                    "column": col,
                    "q1": float(q1),
                    "q3": float(q3),
                    "lower": float(lower),
                    "upper": float(upper),
                    "skipped_zero_iqr": bool(iqr == 0),
                }
            )
            if iqr > 0:
                flag.loc[group.index] = group[col].lt(lower) | group[col].gt(upper)
        data["outlier_" + col] = flag
        data.loc[flag, "outlier_fields"] = data.loc[flag, "outlier_fields"] + col + "|"
        data["is_outlier"] = data["is_outlier"] | flag
        metrics["outliers"]["counts"][col] = int(flag.sum())
    data["outlier_fields"] = data["outlier_fields"].str.rstrip("|")
    metrics["outliers"]["rows_flagged"] = int(data["is_outlier"].sum())
    metrics["scaling"] = {}
    for original, derived in [
        ("temperature_c", "temperature_zscore"),
        ("energy_kwh", "energy_zscore"),
    ]:
        groups = data.groupby("city", observed=True)[original]
        mean = groups.transform("mean")
        std = groups.transform(lambda s: s.std(ddof=0))
        data[derived] = ((data[original] - mean) / std.where(std.ne(0))).fillna(0.0)
        metrics["scaling"][derived] = {
            str(city): {"mean": float(g.mean()), "std_population": float(g.std(ddof=0))}
            for city, g in groups
        }
    data["humidity_fraction"] = data["humidity_pct"] / 100.0
    data["precipitation_log1p"] = np.log1p(data["precipitation_mm"])
    data["timezone"] = TZ
    data["observation_time_utc"] = data["observation_time"].dt.tz_convert("UTC")
    data = data.sort_values(["city", "observation_time"], kind="stable").reset_index(drop=True)
    rejected = (
        pd.concat(rejected_parts, ignore_index=True)
        if rejected_parts
        else pd.DataFrame(columns=REQUIRED_COLUMNS + ["source_row", "rejection_reason"])
    )
    ops["rows_rejected"] = len(rejected)
    ops["rejection_reasons"] = {
        str(k): int(v) for k, v in rejected["rejection_reason"].value_counts().items()
    }
    metrics["after"] = profile(data)
    assert len(raw) == len(data) + len(rejected) + ops["exact_duplicates_removed"]
    assert not data.duplicated(["city", "observation_time"]).any()
    assert not data[REQUIRED_COLUMNS].isna().any().any()
    return data, rejected, metrics


def representative_sample(data: pd.DataFrame, per_group: int = 20, seed: int = 42) -> pd.DataFrame:
    """Toma hasta el límite de filas por ciudad y mes; incluye completos los grupos pequeños."""
    if per_group < 1:
        raise ValueError("per_group debe ser positivo")
    groups = [
        g.sample(n=min(per_group, len(g)), random_state=seed)
        for _, g in data.groupby(["city", "year", "month"], sort=True, observed=True)
    ]
    if not groups:
        return data.iloc[:0].copy()
    return (
        pd.concat(groups)
        .sort_values(["city", "observation_time"], kind="stable")
        .reset_index(drop=True)
    )


def export_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for col in ["observation_time", "observation_time_utc"]:
        if col in result:
            result[col] = result[col].map(
                lambda value: value.isoformat() if pd.notna(value) else ""
            )
    return result


def write_report(metrics: dict, path: Path) -> None:
    """Escribe el informe de auditoría con los conteos y criterios de la ejecución."""
    before, after, ops = metrics["before"], metrics["after"], metrics["operations"]
    lines = [
        "Auditoría de preprocesamiento y limpieza — EA2",
        "Autor: Juan Sebastian Clavijo Correa",
        f"Ejecución UTC: {metrics['execution']['utc']}",
        f"Estado: {metrics['status']}",
        f"Origen: {metrics['source']['database']}",
        f"SHA-256 origen: {metrics['source']['sha256']}",
        "Entorno: base local DuckDB de EA1 y procesamiento con Pandas.",
        "Pandas trabaja en un solo proceso sobre la base que simula el almacenamiento cloud.",
        "Los datos meteorológicos provienen de Open-Meteo. EA1 no fijó el parámetro models.",
        "Las coordenadas corresponden a las ciudades solicitadas. La energía es sintética (semilla 42).",
        "",
        "Comparación antes y después",
        f"Registros antes: {before['rows']}",
        f"Registros después: {after['rows']}",
        f"Duplicados exactos antes: {before['exact_duplicate_rows']}",
        f"Duplicados exactos eliminados: {ops['exact_duplicates_removed']}",
        f"Filas con claves conflictivas rechazadas: {ops['conflicting_key_rows']}",
        f"Total de filas rechazadas: {ops['rows_rejected']}",
        f"Nulos iniciales en columnas de origen: {sum(before['nulls'].values())}",
        f"Nulos finales en columnas de origen: {sum(after['nulls'][col] for col in REQUIRED_COLUMNS)}",
        f"Ecuación: {before['rows']} = {after['rows']} limpias + {ops['exact_duplicates_removed']} duplicadas + {ops['rows_rejected']} rechazadas",
        "",
        "Operaciones y criterios de limpieza",
        "1. Se comprobaron las claves de las dimensiones y los hechos sin correspondencia antes del JOIN.",
        f"2. Normalización de nombres de ciudad: {ops['city_values_normalized']} cambios.",
        f"3. Fechas interpretadas en America/Bogota; fechas inválidas: {ops['invalid_datetimes']}.",
        "4. Se revisaron los tipos de números, identificadores y la marca de energía sintética.",
        f"   Fallos numéricos: {json.dumps(ops['numeric_parse_failures'],ensure_ascii=False)}",
        f"   Valores fuera de dominio convertidos en ausentes: {json.dumps(ops['out_of_domain_values'],ensure_ascii=False)}",
        "5. Los faltantes de temperatura, humedad, viento y energía se completan con la mediana de su ciudad, año y mes.",
        f"   Celdas imputadas: {json.dumps(ops['imputed_cells'],ensure_ascii=False)}",
        "   Una fila sin precipitación válida se rechaza: no permite saber si llovió.",
        "   Si el grupo no tiene valores para calcular la mediana, la fila se rechaza. También se rechazan las versiones que discrepan para una misma ciudad y hora.",
        "6. Año, mes, día, hora, fracción de humedad, log1p de precipitación y z-scores por ciudad.",
        "   Estas variables se añaden a las columnas de origen. El z-score usa la desviación poblacional.",
        "   Si estos datos se usan en un modelo predictivo, las medianas y escalas se deben calcular solo con el conjunto de entrenamiento.",
        "",
        "Valores atípicos",
        metrics["outliers"]["method"],
        metrics["outliers"]["policy"],
        f"Filas con al menos una alerta: {metrics['outliers']['rows_flagged']}",
        f"Alertas por variable: {json.dumps(metrics['outliers']['counts'],ensure_ascii=False)}",
        "Cuando IQR es cero se omite esta clasificación; en series con muchos ceros señalaría cualquier lluvia.",
        "",
        "Tipos y estadísticas por campo",
    ]
    for col in REQUIRED_COLUMNS:
        lines.append(
            f"{col}: {before['dtypes'][col]} -> {after['dtypes'][col]}; nulos {before['nulls'][col]} -> {after['nulls'][col]}"
        )
    for col in MEASURES:
        lines.append(
            f"{col}: ANTES {json.dumps(before['statistics'][col])}; DESPUÉS {json.dumps(after['statistics'][col])}"
        )
    lines += [
        "",
        "Muestra y archivos de la auditoría",
        f"Muestra: {metrics['sampling']['sample_rows']} filas; 20 por ciudad/año/mes como máximo, semilla 42.",
        "La muestra cubre cada ciudad y mes con la misma cuota máxima. Los meses con más registros tienen menor proporción de selección.",
        "Para calcular resultados del conjunto completo, use todas las filas o los pesos N/n de sample_coverage.csv.",
        "Conjunto completo: src/xlsx/cleaned_full.csv",
        "Muestra: src/xlsx/cleaned_data.csv y cleaned_data.xlsx",
        "Rechazos: rejected_records.csv; estadísticas/umbrales: cleaning_metrics.json.",
        "Perfil tabular: quality_profile.csv; cobertura de muestra: sample_coverage.csv.",
        f"Base conservada sin cambios: {metrics['source']['unchanged']}",
        "Los casos con errores se construyen en tests/test_cleaning.py sobre datos separados de EA1.",
        "Un conteo de cero indica que esa operación no encontró filas que corregir en esta ejecución.",
        "",
        "SHA-256 de los archivos generados",
    ]
    lines.extend(f"{name}: {digest}" for name, digest in metrics["output_sha256"].items())
    path.write_text("\n".join(lines) + "\n", encoding="utf8")


def run_pipeline(database: Path, output_root: Path) -> dict:
    """Ejecuta la limpieza y guarda la muestra, el conjunto completo y su auditoría."""
    audit = output_root / "static" / "auditoria"
    out = output_root / "xlsx"
    audit.mkdir(parents=True, exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    raw, source = extract_database(database)
    manifest_path = database.with_name("source_manifest.json")
    if manifest_path.is_file():
        origin = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
        if source["sha256"] != origin["snapshot_sha256"]:
            raise ValueError("La base no coincide con el manifiesto de origen EA1")
    cleaned, rejected, metrics = clean_frame(raw)
    if cleaned.empty:
        raise ValueError("La limpieza no produjo registros válidos; revisar reglas y rechazos")
    sample = representative_sample(cleaned)
    export_frame(cleaned).to_csv(
        out / "cleaned_full.csv", index=False, float_format="%.12g", lineterminator="\n"
    )
    export_frame(sample).to_csv(
        out / "cleaned_data.csv", index=False, float_format="%.12g", lineterminator="\n"
    )
    export_frame(rejected).to_csv(audit / "rejected_records.csv", index=False, lineterminator="\n")
    excel = sample.copy()
    # Excel necesita fechas sin zona horaria. Conservamos las columnas local
    # y UTC por separado; timezone permite interpretar la hora local al abrirlo.
    for col in ["observation_time", "observation_time_utc"]:
        excel[col] = excel[col].dt.tz_localize(None)
    with pd.ExcelWriter(
        out / "cleaned_data.xlsx", engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss"
    ) as writer:
        excel.to_excel(writer, sheet_name="Datos_limpios", index=False)
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        sheet = writer.sheets["Datos_limpios"]
        sheet.freeze_panes = "E2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 42
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, col in enumerate(excel.columns, 1):
            sheet.column_dimensions[get_column_letter(i)].width = (
                26 if "time" in col else max(14, min(30, len(col) + 3))
            )
            for cells in sheet.iter_rows(min_row=2, min_col=i, max_col=i):
                cell = cells[0]
                if col in ["observation_time", "observation_time_utc"]:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif col == "humidity_fraction":
                    cell.number_format = "0.0%"
                elif pd.api.types.is_float_dtype(excel[col]):
                    cell.number_format = "0.0000" if col in ["latitude", "longitude"] else "0.00"
    full_groups = (
        cleaned.groupby(["city", "year", "month"], observed=True).size().rename("population_rows")
    )
    sample_groups = (
        sample.groupby(["city", "year", "month"], observed=True).size().rename("sample_rows")
    )
    coverage = pd.concat([full_groups, sample_groups], axis=1).reset_index()
    coverage["sampling_weight"] = coverage["population_rows"] / coverage["sample_rows"]
    coverage.to_csv(
        audit / "sample_coverage.csv", index=False, float_format="%.8g", lineterminator="\n"
    )
    profile_rows = []
    for stage in ["before", "after"]:
        for col in REQUIRED_COLUMNS:
            profile_rows.append(
                {
                    "stage": stage,
                    "column": col,
                    "dtype": metrics[stage]["dtypes"][col],
                    "nulls": metrics[stage]["nulls"][col],
                    **metrics[stage]["statistics"].get(col, {}),
                }
            )
    pd.DataFrame(profile_rows).to_csv(
        audit / "quality_profile.csv", index=False, lineterminator="\n"
    )
    source.update(
        database="src/db/" + database.name, unchanged=sha256(database) == source["sha256"]
    )
    if not source["unchanged"]:
        raise RuntimeError("La base de origen cambió")
    metrics.update(
        status="PASS",
        source=source,
        execution={
            "utc": datetime.now(timezone.utc).isoformat(),
            "python": sys.version.split()[0],
            "dependencies": {
                p: importlib.metadata.version(p) for p in ["pandas", "numpy", "duckdb", "openpyxl"]
            },
        },
        sampling={
            "sample_rows": len(sample),
            "strata": len(coverage),
            "per_group": 20,
            "seed": 42,
            "design": "Cuota aleatoria fija por ciudad/año/mes; grupos menores se incluyen completos",
        },
        output_sha256={
            file.relative_to(output_root).as_posix(): sha256(file)
            for file in [
                out / "cleaned_full.csv",
                out / "cleaned_data.csv",
                out / "cleaned_data.xlsx",
                audit / "rejected_records.csv",
                audit / "sample_coverage.csv",
                audit / "quality_profile.csv",
            ]
        },
    )
    (audit / "cleaning_metrics.json").write_text(
        json.dumps(metrics, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf8"
    )
    write_report(metrics, audit / "cleaning_report.txt")
    (audit / "cleaning_failure.txt").unlink(missing_ok=True)
    return metrics


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=ROOT / "src" / "db" / "ingestion.db")
    parser.add_argument("--output-root", type=Path, default=ROOT / "src")
    args = parser.parse_args()
    try:
        result = run_pipeline(args.db.resolve(), args.output_root.resolve())
    except Exception as exc:
        audit = args.output_root / "static" / "auditoria"
        audit.mkdir(parents=True, exist_ok=True)
        (audit / "cleaning_failure.txt").write_text(
            f"FAIL\n{type(exc).__name__}: {exc}\n", encoding="utf8"
        )
        raise
    print(
        json.dumps(
            {
                "status": result["status"],
                "before": result["before"]["rows"],
                "after": result["after"]["rows"],
                "duplicates_removed": result["operations"]["exact_duplicates_removed"],
                "rows_rejected": result["operations"]["rows_rejected"],
                "sample_rows": result["sampling"]["sample_rows"],
                "outlier_rows": result["outliers"]["rows_flagged"],
                "source_unchanged": result["source"]["unchanged"],
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
