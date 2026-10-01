"""Integra seis fuentes locales con la salida completa de la Actividad 2."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import platform
import sys
import unicodedata
from datetime import datetime, timezone
from io import StringIO
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parents[1]
TIMEZONE = "America/Bogota"
KEY = ["city", "observation_time"]
SOURCE_CONTRACTS = {
    "geography": ("city_geography.json", ["city_key"], ["country_code", "geoname_id", "elevation_m"]),
    "calendar": ("calendar_2025.xlsx", ["local_date"], ["holiday_name", "is_holiday", "is_weekend"]),
    "solar": ("solar_daily_2025.csv", ["city_key", "local_date"], ["shortwave_radiation_sum_mj_m2", "daylight_seconds", "sunshine_seconds"]),
    "municipalities": ("municipalities.html", ["city_key"], ["municipality_code", "municipality_name", "department_code", "department_name", "divipola_year"]),
    "thermal": ("thermal_parameters.xml", ["city_key"], ["reference_temperature_c", "parameter_scope"]),
    "hours": ("hour_bands.txt", ["hour"], ["day_period"]),
}


class DataQualityError(ValueError):
    """Un contrato de datos impide continuar sin alterar el significado."""


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize_city(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value).strip())
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).casefold().split())


def validate_base(frame: pd.DataFrame) -> None:
    required = KEY + ["source_row", "hour", "month", "year", "temperature_c", "synthetic_energy", "timezone"]
    missing = sorted(set(required) - set(frame.columns))
    if missing:
        raise DataQualityError(f"Faltan columnas en la base: {missing}")
    if frame.empty or frame[KEY].isna().any().any() or frame.duplicated(KEY).any():
        raise DataQualityError("La base debe tener registros y una clave ciudad-hora única, sin nulos.")
    if frame["source_row"].isna().any() or frame["source_row"].duplicated().any():
        raise DataQualityError("source_row debe identificar una sola fila de EA2.")
    if not frame["timezone"].eq(TIMEZONE).all():
        raise DataQualityError("La base no está expresada en America/Bogota.")
    times = pd.to_datetime(frame["observation_time"], utc=True, errors="raise").dt.tz_convert(TIMEZONE)
    for field in ["hour", "month", "year"]:
        if not frame[field].eq(getattr(times.dt, field)).all():
            raise DataQualityError(f"{field} no coincide con la fecha local de observación.")
    values = pd.to_numeric(frame["temperature_c"], errors="raise")
    if not np.isfinite(values).all():
        raise DataQualityError("La temperatura base contiene valores no finitos.")


def load_base(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise FileNotFoundError(f"No existe {path}. Ejecuta python scripts/prepare_base.py.")
    with duckdb.connect(str(path), read_only=True) as connection:
        frame = connection.sql("SELECT * FROM cleaned_data ORDER BY source_row").df()
    # DuckDB entrega TIMESTAMPTZ en la zona de la sesión; el cruce usa la fecha de Colombia.
    frame["observation_time"] = pd.to_datetime(frame["observation_time"], utc=True).dt.tz_convert(TIMEZONE)
    frame["observation_time_utc"] = pd.to_datetime(frame["observation_time_utc"], utc=True)
    validate_base(frame)
    return frame


def verify_source_hashes(directory: Path) -> dict:
    manifest = json.loads((directory / "source_manifest.json").read_text(encoding="utf-8"))
    entries = {Path(entry["path"]).name: entry for entry in manifest["sources"]}
    for filename, _, _ in SOURCE_CONTRACTS.values():
        if filename not in entries:
            raise DataQualityError(f"No hay procedencia registrada para {filename}.")
        if sha256(directory / filename) != entries[filename]["sha256"]:
            raise DataQualityError(f"Cambió {filename}: revisa su procedencia antes de actualizar el manifiesto.")
    return manifest


def load_sources(directory: Path) -> dict[str, pd.DataFrame]:
    frames = {
        "geography": pd.read_json(directory / "city_geography.json"),
        "calendar": pd.read_excel(directory / "calendar_2025.xlsx", sheet_name="calendar", keep_default_na=False),
        "solar": pd.read_csv(directory / "solar_daily_2025.csv", keep_default_na=False, na_values=[""]),
        "municipalities": pd.read_html(
            StringIO((directory / "municipalities.html").read_text(encoding="utf-8")),
            attrs={"id": "divipola"},
            converters={"municipality_code": str, "department_code": str},
        )[0],
        "thermal": pd.read_xml(directory / "thermal_parameters.xml", xpath="./record", parser="etree"),
        "hours": pd.read_csv(directory / "hour_bands.txt", sep="\t", comment="#", keep_default_na=False),
    }
    for name, frame in frames.items():
        if "city" in frame:
            if frame["city"].isna().any() or frame["city"].astype(str).str.strip().eq("").any():
                raise DataQualityError(f"{name}: ciudad vacía.")
            frame["city_key"] = frame["city"].map(normalize_city)
        if "local_date" in frame:
            frame["local_date"] = pd.to_datetime(frame["local_date"], errors="raise").dt.normalize()
        _, keys, fields = SOURCE_CONTRACTS[name]
        missing = set(keys + fields) - set(frame.columns)
        if missing:
            raise DataQualityError(f"{name}: faltan columnas {sorted(missing)}.")
        if frame[keys].isna().any().any() or frame.duplicated(keys).any():
            raise DataQualityError(f"{name}: clave nula o repetida después de normalizar {keys}.")
        frames[name] = frame[keys + fields].copy()
    for field in ["is_holiday", "is_weekend"]:
        if not frames["calendar"][field].isin([True, False]).all():
            raise DataQualityError(f"calendar: {field} requiere booleanos.")
        frames["calendar"][field] = frames["calendar"][field].astype("boolean")
    cities = frames["municipalities"]
    if not cities["municipality_code"].str.fullmatch(r"\d{5}").all() or not cities["department_code"].str.fullmatch(r"\d{2}").all():
        raise DataQualityError("DIVIPOLA requiere códigos de texto de cinco y dos dígitos.")
    if not cities["municipality_code"].str[:2].eq(cities["department_code"]).all():
        raise DataQualityError("El código municipal no corresponde al departamento.")
    if not frames["hours"]["hour"].isin(range(24)).all():
        raise DataQualityError("hour_bands: hora fuera de 0 a 23.")
    numeric_fields = {
        "solar": ["shortwave_radiation_sum_mj_m2", "daylight_seconds", "sunshine_seconds"],
        "thermal": ["reference_temperature_c"],
        "geography": ["geoname_id", "elevation_m"],
    }
    for name, fields in numeric_fields.items():
        for field in fields:
            series = pd.to_numeric(frames[name][field], errors="raise")
            if np.isinf(series.dropna()).any():
                raise DataQualityError(f"{name}: infinito en {field}.")
            if name == "solar" and (series.dropna() < 0).any():
                raise DataQualityError(f"solar: valor negativo en {field}.")
            frames[name][field] = series
    solar = frames["solar"]
    if (solar["sunshine_seconds"] > solar["daylight_seconds"]).any() or (solar["daylight_seconds"] > 86400).any():
        raise DataQualityError("Las duraciones solares no son coherentes.")
    return frames


def integrate(base: pd.DataFrame, sources: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, list[dict]]:
    validate_base(base)
    enriched = base.copy()
    times = pd.to_datetime(enriched["observation_time"], utc=True).dt.tz_convert(TIMEZONE)
    enriched["local_date"] = times.dt.tz_localize(None).dt.normalize()
    enriched["city_key"] = enriched["city"].map(normalize_city)
    audit = []
    for name, (filename, keys, fields) in SOURCE_CONTRACTS.items():
        source = sources[name]
        if source[keys].isna().any().any() or source.duplicated(keys).any():
            raise DataQualityError(f"{name}: el cruce exige una fila por {keys}.")
        collisions = set(fields) & set(enriched.columns)
        if collisions:
            raise DataQualityError(f"{name}: se intentó sobrescribir {sorted(collisions)}.")
        keys_used = enriched[keys].drop_duplicates()
        unused = source[keys].merge(keys_used, on=keys, how="left", indicator=True)
        before = len(enriched)
        enriched = enriched.merge(source, on=keys, how="left", sort=False, validate="many_to_one", indicator="_join")
        matches = enriched["_join"].eq("both")
        enriched[f"matched_{name}"] = matches
        audit.append({
            "source": name, "file": filename, "join": "left many_to_one", "keys": keys,
            "source_rows": len(source), "before_rows": before, "after_rows": len(enriched),
            "matched_rows": int(matches.sum()), "unmatched_rows": int((~matches).sum()),
            "unmatched_keys": enriched.loc[~matches, keys].drop_duplicates().astype(str).head(10).to_dict("records"),
            "unused_source_keys": int(unused["_merge"].eq("left_only").sum()),
            "payload_null_cells": int(enriched[fields].isna().sum().sum()),
        })
        enriched = enriched.drop(columns="_join")
        if len(enriched) != before:
            raise DataQualityError(f"{name}: el cruce alteró el número de registros.")
    enriched = enriched.sort_values("source_row", kind="stable").reset_index(drop=True)
    original = base.sort_values("source_row", kind="stable").reset_index(drop=True)
    pd.testing.assert_frame_equal(enriched[base.columns], original, check_exact=True)
    enriched["radiation_kwh_m2_day"] = enriched["shortwave_radiation_sum_mj_m2"] / 3.6
    enriched["daylight_hours"] = enriched["daylight_seconds"] / 3600
    enriched["sunshine_hours"] = enriched["sunshine_seconds"] / 3600
    enriched["degrees_above_reference_c"] = (enriched["temperature_c"] - enriched["reference_temperature_c"]).clip(lower=0)
    enriched["degrees_below_reference_c"] = (enriched["reference_temperature_c"] - enriched["temperature_c"]).clip(lower=0)
    enriched["is_business_day"] = ~(enriched["is_holiday"].astype("boolean") | enriched["is_weekend"].astype("boolean"))
    payload = [field for _, _, fields in SOURCE_CONTRACTS.values() for field in fields]
    complete = enriched[[f"matched_{name}" for name in SOURCE_CONTRACTS]].all(axis=1) & enriched[payload].notna().all(axis=1)
    enriched["enrichment_status"] = np.where(complete, "COMPLETE", "INCOMPLETE")
    return enriched.drop(columns="city_key"), audit


def representative_sample(frame: pd.DataFrame, quota: int = 20, seed: int = 42) -> tuple[pd.DataFrame, pd.DataFrame]:
    if quota < 1:
        raise ValueError("La cuota de muestra debe ser positiva.")
    pieces, coverage = [], []
    for (city, year, month), group in frame.groupby(["city", "year", "month"], sort=True):
        n = min(quota, len(group))
        pieces.append(group.sample(n=n, random_state=seed))
        coverage.append({"city": city, "year": int(year), "month": int(month), "population_rows": len(group), "sample_rows": n, "sample_weight": len(group) / n})
    sample = pd.concat(pieces).sort_values("source_row", kind="stable").reset_index(drop=True)
    return sample, pd.DataFrame(coverage)


def serializable_frame(frame: pd.DataFrame, excel: bool = False) -> pd.DataFrame:
    result = frame.copy()
    for field in ["observation_time", "observation_time_utc"]:
        if excel:
            result[field] = result[field].dt.tz_localize(None)
        else:
            result[field] = result[field].map(lambda value: value.isoformat())
    if not excel:
        result["local_date"] = result["local_date"].dt.strftime("%Y-%m-%d")
    return result


def write_workbook(sample: pd.DataFrame, coverage: pd.DataFrame, path: Path, base_columns: list[str]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        serializable_frame(sample, excel=True).to_excel(writer, sheet_name="Datos", index=False)
        coverage.to_excel(writer, sheet_name="Cobertura", index=False, startrow=4)
        data = writer.sheets["Datos"]
        data.sheet_view.showGridLines = False
        data.freeze_panes = "D2"
        data.auto_filter.ref = data.dimensions
        data.row_dimensions[1].height = 58
        for index, field in enumerate(sample.columns, start=1):
            letter = get_column_letter(index)
            width = 25 if field in ["observation_time", "observation_time_utc"] else 22
            if field in ["outlier_fields", "holiday_name", "municipality_name", "department_name"]:
                width = 34
            if field in ["degrees_above_reference_c", "degrees_below_reference_c"]:
                width = 32
            elif field == "reference_temperature_c":
                width = 28
            data.column_dimensions[letter].width = width
            cell = data.cell(1, index)
            cell.fill = PatternFill("solid", fgColor="17365D" if field in base_columns else "176B72")
            cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            for row in range(2, len(sample) + 2):
                cell = data.cell(row, index)
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top")
                if field in ["observation_time", "observation_time_utc"]:
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif field == "local_date":
                    cell.number_format = "yyyy-mm-dd"
                elif field in ["municipality_code", "department_code"]:
                    cell.number_format = "@"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0000"
                if field in ["outlier_fields", "holiday_name"]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in range(2, len(sample) + 2):
            record = sample.iloc[row - 2]
            lines = max(math.ceil(len(str(record[field])) / 34) for field in ["outlier_fields", "holiday_name"])
            data.row_dimensions[row].height = max(20, lines * 16 + 4)
        cover = writer.sheets["Cobertura"]
        cover.sheet_view.showGridLines = False
        cover["A1"] = "Muestra estratificada por ciudad, año y mes"
        cover["A2"] = "20 filas por grupo, semilla 42. El peso permite reconstruir el tamaño de cada estrato."
        cover["A3"] = "Energía sintética. Fechas locales sin zona en Excel; timezone conserva America/Bogota."
        cover["A4"] = "Procedencia y limitaciones: ../sources/source_manifest.json y ../../docs/FUENTES.md."
        for row in range(1, 5):
            cover.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cover.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
            cover.row_dimensions[row].height = 30
        cover["A1"].font = Font(name="Calibri", bold=True, size=14, color="17365D")
        for cell in cover[5]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cover.column_dimensions[cell.column_letter].width = 21
        cover.freeze_panes = "A6"
        cover.auto_filter.ref = f"A5:F{5 + len(coverage)}"
        for row in range(6, 6 + len(coverage)):
            cover.cell(row, 6).number_format = "0.00"


def run(database: Path, source_directory: Path, output: Path) -> dict:
    input_hash = sha256(database)
    base_manifest_path = database.with_name("base_manifest.json")
    base_manifest = json.loads(base_manifest_path.read_text(encoding="utf-8")) if base_manifest_path.exists() else {}
    if base_manifest and base_manifest["database"]["sha256"] != input_hash:
        raise DataQualityError("La base no coincide con la instantánea de EA2 registrada en base_manifest.json.")
    manifest = verify_source_hashes(source_directory)
    base = load_base(database)
    sources = load_sources(source_directory)
    enriched, joins = integrate(base, sources)
    sample, coverage = representative_sample(enriched)
    for relative in ["xlsx", "data", "static/auditoria"]:
        (output / relative).mkdir(parents=True, exist_ok=True)
    audit_directory = output / "static/auditoria"
    full_path = output / "data/enriched_full.csv"
    sample_path = output / "xlsx/enriched_data.csv"
    excel_path = output / "xlsx/enriched_data.xlsx"
    serializable_frame(enriched).to_csv(full_path, index=False, lineterminator="\n")
    serializable_frame(sample).to_csv(sample_path, index=False, lineterminator="\n")
    coverage.to_csv(audit_directory / "sample_coverage.csv", index=False, lineterminator="\n")
    write_workbook(sample, coverage, excel_path, list(base.columns))
    if sha256(database) != input_hash:
        raise DataQualityError("La base de entrada cambió durante la ejecución.")
    summary = {
        "execution_utc": datetime.now(timezone.utc).isoformat(), "python": platform.python_version(),
        "dependencies": {name: importlib.metadata.version(name) for name in ["pandas", "numpy", "duckdb", "openpyxl", "lxml", "tzdata"]},
        "upstream_ea2": base_manifest.get("source", {}),
        "rows_before": len(base), "rows_after": len(enriched), "columns_before": len(base.columns),
        "columns_after": len(enriched.columns), "sample_rows": len(sample), "sample_strata": len(coverage),
        "original_columns_preserved": list(base.columns), "original_values_unchanged": True,
        "input_database_sha256": input_hash, "input_database_unchanged": True,
        "duplicate_city_hour_after": int(enriched.duplicated(KEY).sum()),
        "incomplete_rows": int(enriched["enrichment_status"].eq("INCOMPLETE").sum()),
        "synthetic_energy_rows": int(base["synthetic_energy"].sum()),
        "joins": joins, "output_columns": list(enriched.columns),
        "outputs_sha256": {path.relative_to(output).as_posix(): sha256(path) for path in [full_path, sample_path, excel_path]},
    }
    (audit_directory / "integration_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (audit_directory / "source_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [
        "AUDITORÍA DEL ENRIQUECIMIENTO", f"Ejecución UTC: {summary['execution_utc']}",
        f"Registros antes: {len(base):,}. Registros después: {len(enriched):,}.",
        f"Columnas: {len(base.columns)} -> {len(enriched.columns)}. Valores originales conservados: sí.",
        f"Base DuckDB leída sin escritura. SHA-256: {input_hash}",
        "Clave del conjunto: ciudad y hora local America/Bogota.",
        f"CSV original EA2 SHA-256: {base_manifest.get('source', {}).get('sha256', 'No registrado para esta entrada personalizada')}",
        f"Commit EA2: {base_manifest.get('source', {}).get('commit', 'No registrado')}",
        "", "CRUCES POR FUENTE",
    ]
    provenance = {entry["path"]: entry for entry in manifest["sources"]}
    for join in joins:
        lines.append(f"{join['file']}: LEFT JOIN many_to_one por {', '.join(join['keys'])}; fuente={join['source_rows']}, coincidentes={join['matched_rows']}, sin correspondencia={join['unmatched_rows']}, claves fuente sin uso={join['unused_source_keys']}, celdas adicionales nulas={join['payload_null_cells']}. Filas {join['before_rows']} -> {join['after_rows']}.")
        notes = provenance.get(join["file"], {}).get("notes", [])
        if notes:
            lines.append("Observaciones: " + " ".join(notes))
        if join["unmatched_keys"]:
            lines.append(f"Ejemplos sin correspondencia: {join['unmatched_keys']}")
    lines.extend([
        "", "TRANSFORMACIONES Y CONTROLES",
        "Ciudades: se ignoran tildes, mayúsculas y espacios únicamente en la clave auxiliar. city original se conserva.",
        "Fechas: cruce diario según America/Bogota, sin trasladar las horas nocturnas al día UTC siguiente.",
        "DIVIPOLA: códigos tratados como texto para conservar ceros iniciales; city_id es una clave local de EA1.",
        "Unidades: MJ/m²/día divididos por 3,6 = kWh/m²/día; segundos solares divididos por 3.600 = horas.",
        "Diferencias térmicas: máximo(temperatura - referencia, 0) y máximo(referencia - temperatura, 0).",
        "Referencia térmica y franjas horarias: parámetros académicos, no mediciones, tarifas ni normas de confort.",
        "No se eliminan filas ni se imputan mediciones adicionales. Una clave de fuente repetida detiene la ejecución.",
        f"Filas incompletas: {summary['incomplete_rows']}; se preservan y señalan con enrichment_status.",
        f"Muestra: {len(sample)} filas, {len(coverage)} estratos ciudad-año-mes, cuota máxima 20, semilla 42.",
        "La cuota no es una muestra aleatoria simple global. sample_coverage.csv incluye los pesos por estrato.",
        "", "INTERPRETACIÓN",
        "energy_kwh es sintética en toda la base. No se midió ahorro ni mejora predictiva en esta actividad.",
        "La radiación diaria describe el día completo: no usarla para anticipar una hora de ese mismo día.",
        "Antes de modelar, dividir temporalmente y recalcular escalados de EA2 usando solo entrenamiento.",
        "Las fuentes meteorológicas son productos de modelo/reanálisis, no sensores municipales.",
        "Consultar FUENTES.md: instantáneas, versiones, año del registro DIVIPOLA y licencias.",
    ])
    (audit_directory / "enriched_report.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Enriquecimiento: {len(base):,} -> {len(enriched):,} registros; {len(enriched.columns)} columnas; muestra {len(sample):,}.")
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=PROJECT / "src/db/cleaned.duckdb")
    parser.add_argument("--sources", type=Path, default=PROJECT / "src/sources")
    parser.add_argument("--output-dir", type=Path, default=PROJECT / "src")
    args = parser.parse_args()
    try:
        run(args.db, args.sources, args.output_dir)
    except (DataQualityError, FileNotFoundError, ValueError, duckdb.Error) as error:
        print(f"No se completó el enriquecimiento: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
