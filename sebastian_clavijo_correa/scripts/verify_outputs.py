"""Compara las salidas con el CSV original sin volver a ejecutar los cruces."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def verify() -> dict:
    report = json.loads(
        (SRC / "static/auditoria/integration_summary.json").read_text(encoding="utf-8")
    )
    manifest = json.loads((SRC / "db/base_manifest.json").read_text(encoding="utf-8"))
    base_path = ROOT / manifest["csv"]["path"]
    require(
        digest(base_path) == manifest["source"]["sha256"], "El CSV ya no es la copia exacta de EA2."
    )
    require(
        digest(SRC / "db/cleaned.duckdb") == manifest["database"]["sha256"],
        "Cambió la base analítica.",
    )
    base = pd.read_csv(base_path, dtype=str, keep_default_na=False)
    full = pd.read_csv(SRC / "data/enriched_full.csv", dtype=str, keep_default_na=False)
    sample = pd.read_csv(SRC / "xlsx/enriched_data.csv", dtype=str, keep_default_na=False)
    require(
        len(full) == len(base) == 43800, "El conjunto completo debe conservar las 43.800 filas."
    )
    require(len(full.columns) == 63, "El contrato de esta versión exige 63 columnas.")
    require(
        not full.duplicated(["city", "observation_time"]).any(), "Se duplicó la clave ciudad-hora."
    )
    require(
        full["source_row"].tolist() == base["source_row"].tolist(),
        "Cambió el orden o la identidad de las filas.",
    )
    schema = {row["column"]: row["type"] for row in manifest["database"]["schema"]}
    for field in base.columns:
        if field in ["observation_time", "observation_time_utc"]:
            left = pd.to_datetime(base[field], utc=True)
            right = pd.to_datetime(full[field], utc=True)
            require(left.eq(right).all(), f"Cambió el instante en {field}.")
        elif schema[field] in ["DOUBLE", "BIGINT"]:
            require(
                np.array_equal(base[field].map(float), full[field].map(float)),
                f"Cambió el valor de {field}.",
            )
        else:
            require(base[field].eq(full[field]).all(), f"Cambió el contenido de {field}.")
    require(
        len(sample) == 1200 and not sample["source_row"].duplicated().any(),
        "La muestra no tiene 1.200 filas distintas.",
    )
    selected = (
        full.set_index("source_row", drop=False).loc[sample["source_row"]].reset_index(drop=True)
    )
    pd.testing.assert_frame_equal(selected, sample, check_exact=True)
    groups = sample.groupby(["city", "year", "month"]).size()
    require(
        len(groups) == 60 and groups.eq(20).all(),
        "La muestra no cubre los 60 estratos con 20 filas.",
    )
    coverage = pd.read_csv(SRC / "static/auditoria/sample_coverage.csv")
    require(
        coverage["population_rows"].sum() == len(full),
        "La cobertura no reconcilia con la población.",
    )
    require(
        coverage["sample_rows"].sum() == len(sample), "La cobertura no reconcilia con la muestra."
    )
    require(
        np.allclose(
            coverage["sample_weight"] * coverage["sample_rows"], coverage["population_rows"]
        ),
        "Los pesos no corresponden a los estratos.",
    )
    require(full["synthetic_energy"].eq("True").all(), "Se perdió la marca de energía sintética.")
    require(
        full["enrichment_status"].eq("COMPLETE").all(),
        "Esta entrega contiene registros incompletos; revisar auditoría.",
    )
    require(
        full["municipality_code"].str.fullmatch(r"\d{5}").all(), "Se perdieron dígitos DIVIPOLA."
    )
    require(
        full.loc[full["city"].eq("Medellin"), "municipality_code"].eq("05001").all(),
        "Medellín perdió el cero inicial.",
    )
    local_dates = (
        pd.to_datetime(full["observation_time"], utc=True)
        .dt.tz_convert("America/Bogota")
        .dt.strftime("%Y-%m-%d")
    )
    require(local_dates.eq(full["local_date"]).all(), "El cruce diario no conserva la fecha local.")
    require(
        np.allclose(
            full["radiation_kwh_m2_day"].astype(float) * 3.6,
            full["shortwave_radiation_sum_mj_m2"].astype(float),
        ),
        "Error en conversión de radiación.",
    )
    require(
        np.allclose(
            full["daylight_hours"].astype(float) * 3600, full["daylight_seconds"].astype(float)
        ),
        "Error en conversión de duración.",
    )
    excel_path = SRC / "xlsx/enriched_data.xlsx"
    excel = pd.read_excel(
        excel_path,
        sheet_name="Datos",
        keep_default_na=False,
        dtype={"municipality_code": str, "department_code": str},
    )
    require(excel.shape == sample.shape, "El Excel y el CSV tienen dimensiones diferentes.")
    require(
        excel.columns.tolist() == sample.columns.tolist(), "Cambió el orden de columnas en Excel."
    )
    for field in sample.columns:
        if field in ["observation_time", "observation_time_utc"]:
            expected = pd.to_datetime(sample[field], utc=True)
            if field == "observation_time":
                expected = expected.dt.tz_convert("America/Bogota")
            require(
                pd.to_datetime(excel[field]).eq(expected.dt.tz_localize(None)).all(),
                f"Excel cambió {field}.",
            )
        elif field == "local_date":
            require(
                pd.to_datetime(excel[field]).dt.strftime("%Y-%m-%d").eq(sample[field]).all(),
                "Excel cambió la fecha local.",
            )
        elif pd.api.types.is_numeric_dtype(excel[field]) and not pd.api.types.is_bool_dtype(
            excel[field]
        ):
            require(
                np.allclose(
                    excel[field].to_numpy(dtype=float),
                    sample[field].to_numpy(dtype=float),
                    rtol=1e-12,
                    atol=1e-12,
                ),
                f"Excel cambió números en {field}.",
            )
        else:
            require(
                excel[field].astype(str).eq(sample[field]).all(),
                f"Excel cambió textos/booleanos en {field}.",
            )
    workbook = load_workbook(excel_path, read_only=False, data_only=False)
    require(workbook.sheetnames == ["Datos", "Cobertura"], "Hojas inesperadas en la evidencia.")
    require(workbook["Datos"].freeze_panes == "D2", "Faltan paneles de identificación.")
    require(bool(workbook["Datos"].auto_filter.ref), "Falta el filtro de la muestra.")
    require(
        not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row),
        "La evidencia debe contener valores exportados, no fórmulas.",
    )
    workbook.close()
    for path, expected in report["outputs_sha256"].items():
        require(
            digest(SRC / path.replace("\\", "/")) == expected,
            f"Cambió una salida después de la auditoría: {path}",
        )
    require(
        report["rows_before"] == report["rows_after"] == len(full),
        "La auditoría no concuerda con los datos.",
    )
    for join in report["joins"]:
        require(
            join["matched_rows"] == len(full) and join["unmatched_rows"] == 0,
            f"Cobertura incompleta: {join['source']}",
        )
    require(
        (ROOT / ".github/workflows/bigdata.yml").read_bytes()
        == (ROOT.parent / ".github/workflows/bigdata.yml").read_bytes(),
        "Las copias del workflow difieren.",
    )
    result = {
        "status": "PASS",
        "base_rows": len(base),
        "enriched_rows": len(full),
        "columns": len(full.columns),
        "source_cells_verified": len(base) * len(base.columns),
        "sample_rows": len(sample),
        "sample_cells_compared_with_excel": len(sample) * len(sample.columns),
        "strata": len(groups),
        "city_hour_duplicates": 0,
        "original_values_preserved": True,
        "xlsx_csv_equivalent": True,
    }
    (SRC / "static/auditoria/verification.json").write_text(
        json.dumps(result, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False))
    return result


if __name__ == "__main__":
    verify()
