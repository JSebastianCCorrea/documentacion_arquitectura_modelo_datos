"""Pruebas de integración de fuentes y conservación de los datos de EA2.

Los casos defectuosos se construyen en carpetas temporales. Ninguna prueba
consulta una API ni escribe en la base o en las fuentes del proyecto.
"""

from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from xml.etree import ElementTree as ET

import duckdb
import numpy as np
import pandas as pd
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "src"))
import enrichment as pipeline


def example_base(times=None, temperatures=None):
    """Tres observaciones bastan para distinguir umbral, unidades y fecha local."""
    if times is None:
        times = ["2025-01-01 00:00", "2025-01-01 01:00", "2025-01-01 23:00"]
    local = pd.DatetimeIndex(times).tz_localize(pipeline.TIMEZONE)
    n = len(local)
    if temperatures is None:
        temperatures = [15.0, 18.0, 21.0][:n]
    return pd.DataFrame({
        "city_id": [5] * n,
        "time_id": np.arange(1, n + 1),
        "observation_time": local,
        "city": ["Medellín"] * n,
        "latitude": [6.2518] * n,
        "longitude": [-75.5636] * n,
        "temperature_c": temperatures,
        "humidity_pct": [65.0] * n,
        "precipitation_mm": [0.0] * n,
        "wind_speed_kmh": [4.0] * n,
        "energy_kwh": [123.45] * n,
        "synthetic_energy": [True] * n,
        "source_row": np.arange(1, n + 1),
        "year": local.year,
        "month": local.month,
        "day": local.day,
        "hour": local.hour,
        "imputed_fields": [""] * n,
        "quality_status": ["VALID"] * n,
        "is_outlier": [False] * n,
        "outlier_fields": [""] * n,
        "outlier_temperature_c": [False] * n,
        "outlier_humidity_pct": [False] * n,
        "outlier_precipitation_mm": [False] * n,
        "outlier_wind_speed_kmh": [False] * n,
        "outlier_energy_kwh": [False] * n,
        "temperature_zscore": [-0.2] * n,
        "energy_zscore": [0.1] * n,
        "humidity_fraction": [0.65] * n,
        "precipitation_log1p": [0.0] * n,
        "timezone": [pipeline.TIMEZONE] * n,
        "observation_time_utc": local.tz_convert("UTC"),
    })


def make_sources(directory):
    """Fuentes mínimas en los seis formatos, sin usar la red."""
    directory.mkdir(parents=True, exist_ok=True)
    geography = [{"city": "  MEDELLÍN ", "country_code": "CO", "geoname_id": 3674962, "elevation_m": 1495.0}]
    (directory / "city_geography.json").write_text(json.dumps(geography, ensure_ascii=False), encoding="utf-8")
    dates = pd.date_range("2025-01-01", periods=3)
    calendar = pd.DataFrame({
        "local_date": dates,
        "holiday_name": ["Año Nuevo", "No festivo", "No festivo"],
        "is_holiday": [True, False, False],
        "is_weekend": [False, False, False],
    })
    calendar.to_excel(directory / "calendar_2025.xlsx", sheet_name="calendar", index=False)
    solar = pd.DataFrame({
        "city": ["Medellin"] * 3,
        "local_date": dates.strftime("%Y-%m-%d"),
        "shortwave_radiation_sum_mj_m2": [7.2, 14.4, 0.0],
        "daylight_seconds": [43200.0, 43200.0, 43200.0],
        "sunshine_seconds": [10800.0, 7200.0, 0.0],
    })
    solar.to_csv(directory / "solar_daily_2025.csv", index=False)
    municipalities = pd.DataFrame([{
        "city": "Medellín", "municipality_code": "05001", "municipality_name": "MEDELLÍN",
        "department_code": "05", "department_name": "ANTIOQUIA", "divipola_year": 2025,
    }])
    (directory / "municipalities.html").write_text(municipalities.to_html(index=False, table_id="divipola"), encoding="utf-8")
    root = ET.Element("parameters")
    record = ET.SubElement(root, "record")
    for key, value in {"city": "Medellín", "reference_temperature_c": "18.0", "parameter_scope": "academic_illustration"}.items():
        ET.SubElement(record, key).text = value
    ET.ElementTree(root).write(directory / "thermal_parameters.xml", encoding="utf-8", xml_declaration=True)
    pd.DataFrame({
        "hour": range(24),
        "day_period": ["madrugada"] * 6 + ["mañana"] * 6 + ["tarde"] * 6 + ["noche"] * 6,
    }).to_csv(directory / "hour_bands.txt", sep="\t", index=False)
    manifest = {
        "period": {"start": "2025-01-01", "end": "2025-12-31", "timezone": pipeline.TIMEZONE},
        "network_required_for_pipeline": False,
        "sources": [
            {"path": filename, "sha256": pipeline.sha256(directory / filename)}
            for filename, _, _ in pipeline.SOURCE_CONTRACTS.values()
        ],
    }
    (directory / "source_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return manifest


def write_database(frame, path):
    with duckdb.connect(str(path)) as connection:
        connection.register("fixture", frame)
        connection.execute("CREATE TABLE cleaned_data AS SELECT * FROM fixture")


def assert_exports_agree(test, output):
    """Compara valores de la muestra, incluidos fechas, booleanos y códigos."""
    csv = pd.read_csv(output / "xlsx/enriched_data.csv", dtype=str, keep_default_na=False)
    workbook = load_workbook(output / "xlsx/enriched_data.xlsx", read_only=True, data_only=True)
    try:
        sheet = workbook["Datos"]
        rows = list(sheet.iter_rows(values_only=True))
        test.assertEqual(list(rows[0]), list(csv.columns))
        test.assertEqual(len(rows) - 1, len(csv))
        for index, row in enumerate(rows[1:]):
            for field, value in zip(csv.columns, row):
                text = csv.iloc[index][field]
                with test.subTest(row=index, field=field):
                    if field in {"observation_time", "observation_time_utc"}:
                        expected = pd.Timestamp(text).tz_localize(None).to_pydatetime()
                        test.assertEqual(value, expected)
                    elif field == "local_date":
                        test.assertEqual(value.date(), pd.Timestamp(text).date())
                    elif isinstance(value, bool):
                        test.assertEqual(text.lower(), str(value).lower())
                    elif isinstance(value, (int, float)):
                        test.assertAlmostEqual(float(text), value, delta=max(1e-10, abs(value) * 1e-12))
                    else:
                        test.assertEqual("" if value is None else str(value), text)
        test.assertIn("Cobertura", workbook.sheetnames)
    finally:
        workbook.close()
    return csv


class SourceAndJoinTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory()
        cls.source_dir = Path(cls.folder.name)
        make_sources(cls.source_dir)
        cls.valid_sources = pipeline.load_sources(cls.source_dir)

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def setUp(self):
        self.base = example_base()
        self.sources = {name: frame.copy(deep=True) for name, frame in self.valid_sources.items()}

    def test_city_normalization_handles_accents_case_and_spacing(self):
        self.assertEqual(pipeline.normalize_city(" \t BÓGOTÁ   D.C. \n"), "bogota d.c.")
        self.assertEqual(pipeline.normalize_city("MEDELLÍN"), pipeline.normalize_city("  Medellin "))

    def test_reads_six_formats_and_preserves_leading_zero_codes(self):
        self.assertEqual(set(self.sources), set(pipeline.SOURCE_CONTRACTS))
        self.assertEqual([len(self.sources[name]) for name in pipeline.SOURCE_CONTRACTS], [1, 3, 3, 1, 1, 24])
        self.assertEqual(self.sources["municipalities"].iloc[0]["municipality_code"], "05001")
        self.assertEqual(self.sources["municipalities"].iloc[0]["department_code"], "05")
        self.assertEqual(self.sources["geography"].iloc[0]["city_key"], "medellin")
        self.assertTrue(pd.api.types.is_datetime64_any_dtype(self.sources["calendar"]["local_date"]))
        self.assertTrue(pd.api.types.is_bool_dtype(self.sources["calendar"]["is_holiday"]))

    def test_original_32_columns_and_source_frames_remain_unchanged(self):
        original = self.base.copy(deep=True)
        sources_before = copy.deepcopy(self.sources)
        enriched, audit = pipeline.integrate(self.base.iloc[::-1], self.sources)
        self.assertEqual(len(original.columns), 32)
        pd.testing.assert_frame_equal(enriched[original.columns], original, check_exact=True)
        pd.testing.assert_frame_equal(self.base, original, check_exact=True)
        for name in self.sources:
            pd.testing.assert_frame_equal(self.sources[name], sources_before[name], check_exact=True)
        self.assertEqual(len(audit), 6)
        self.assertTrue(enriched["enrichment_status"].eq("COMPLETE").all())

    def test_duplicate_keys_in_each_source_stop_many_to_one_join(self):
        for name in self.sources:
            with self.subTest(source=name):
                sources = copy.deepcopy(self.sources)
                sources[name] = pd.concat([sources[name], sources[name].iloc[[0]]], ignore_index=True)
                with self.assertRaises(pipeline.DataQualityError):
                    pipeline.integrate(self.base, sources)

    def test_missing_coverage_preserves_rows_and_marks_incomplete(self):
        self.sources["solar"] = self.sources["solar"].iloc[1:].copy()
        enriched, audit = pipeline.integrate(self.base, self.sources)
        self.assertEqual(len(enriched), len(self.base))
        pd.testing.assert_frame_equal(enriched[self.base.columns], self.base, check_exact=True)
        self.assertFalse(enriched["matched_solar"].any())
        self.assertTrue(enriched["enrichment_status"].eq("INCOMPLETE").all())
        self.assertTrue(enriched["radiation_kwh_m2_day"].isna().all())
        self.assertEqual(next(item for item in audit if item["source"] == "solar")["unmatched_rows"], 3)

    def test_matched_source_with_missing_payload_is_still_incomplete(self):
        self.sources["geography"].loc[0, "elevation_m"] = np.nan
        enriched, audit = pipeline.integrate(self.base, self.sources)
        self.assertTrue(enriched["matched_geography"].all())
        self.assertTrue(enriched["enrichment_status"].eq("INCOMPLETE").all())
        self.assertEqual(next(item for item in audit if item["source"] == "geography")["payload_null_cells"], 3)

    def test_late_local_hour_uses_local_date_even_when_utc_is_next_day(self):
        enriched, _ = pipeline.integrate(self.base, self.sources)
        last = enriched.iloc[-1]
        self.assertEqual(last["observation_time_utc"].day, 2)
        self.assertEqual(last["local_date"], pd.Timestamp("2025-01-01"))
        self.assertEqual(last["holiday_name"], "Año Nuevo")
        self.assertTrue(last["is_holiday"])
        self.assertFalse(last["is_business_day"])
        self.assertEqual(last["day_period"], "noche")
        self.assertEqual(last["radiation_kwh_m2_day"], 2.0)

    def test_solar_units_convert_mj_and_seconds(self):
        enriched, _ = pipeline.integrate(self.base, self.sources)
        np.testing.assert_allclose(enriched["radiation_kwh_m2_day"], [2.0] * 3)
        np.testing.assert_allclose(enriched["daylight_hours"], [12.0] * 3)
        np.testing.assert_allclose(enriched["sunshine_hours"], [3.0] * 3)

    def test_thermal_reference_18_has_distinct_below_equal_and_above_cases(self):
        enriched, _ = pipeline.integrate(self.base, self.sources)
        np.testing.assert_array_equal(enriched["degrees_above_reference_c"], [0.0, 0.0, 3.0])
        np.testing.assert_array_equal(enriched["degrees_below_reference_c"], [3.0, 0.0, 0.0])
        self.assertTrue(enriched["parameter_scope"].eq("academic_illustration").all())

    def test_payload_collision_cannot_overwrite_base_column(self):
        self.base["country_code"] = "ORIGINAL"
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.integrate(self.base, self.sources)
        self.assertTrue(self.base["country_code"].eq("ORIGINAL").all())

    def test_duplicate_base_city_hour_is_rejected(self):
        duplicate = self.base.iloc[[0]].copy()
        duplicate["source_row"] = 999
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.integrate(pd.concat([self.base, duplicate], ignore_index=True), self.sources)

    def test_duplicate_source_row_is_rejected(self):
        self.base.loc[1, "source_row"] = self.base.loc[0, "source_row"]
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.validate_base(self.base)

    def test_invalid_base_temperatures_are_rejected(self):
        for invalid in ["no es un número", np.inf, np.nan]:
            with self.subTest(value=invalid):
                base = self.base.copy()
                base["temperature_c"] = base["temperature_c"].astype(object)
                base.loc[0, "temperature_c"] = invalid
                with self.assertRaises(ValueError):
                    pipeline.validate_base(base)

    def test_calendar_fields_must_agree_with_local_time(self):
        for field, value in [("hour", 9), ("month", 12), ("year", 2024), ("timezone", "UTC")]:
            with self.subTest(field=field):
                base = self.base.copy()
                base.loc[0, field] = value
                with self.assertRaises(pipeline.DataQualityError):
                    pipeline.validate_base(base)


class SourceFileValidationTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.directory = Path(self.folder.name)
        self.manifest = make_sources(self.directory)

    def test_normalization_cannot_hide_duplicate_city_keys(self):
        path = self.directory / "city_geography.json"
        records = json.loads(path.read_text(encoding="utf-8"))
        repeated = dict(records[0], city="medellin")
        path.write_text(json.dumps(records + [repeated]), encoding="utf-8")
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.load_sources(self.directory)

    def test_invalid_solar_values_and_durations_are_rejected(self):
        path = self.directory / "solar_daily_2025.csv"
        original = pd.read_csv(path)
        for field, invalid in [
            ("shortwave_radiation_sum_mj_m2", "incorrecto"),
            ("shortwave_radiation_sum_mj_m2", np.inf),
            ("shortwave_radiation_sum_mj_m2", -1.0),
            ("daylight_seconds", 86401.0),
            ("sunshine_seconds", 50000.0),
        ]:
            with self.subTest(field=field, value=invalid):
                frame = original.copy()
                frame[field] = frame[field].astype(object)
                frame.loc[0, field] = invalid
                frame.to_csv(path, index=False)
                with self.assertRaises(ValueError):
                    pipeline.load_sources(self.directory)

    def test_hash_mismatch_stops_run_before_writing_outputs(self):
        pipeline.verify_source_hashes(self.directory)
        path = self.directory / "hour_bands.txt"
        path.write_text(path.read_text(encoding="utf-8") + "# alterado\n", encoding="utf-8")
        database = self.directory / "fixture.db"
        write_database(example_base(), database)
        output = self.directory / "outputs"
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.run(database, self.directory, output)
        self.assertFalse(output.exists())

    def test_manifest_must_identify_every_required_source(self):
        self.manifest["sources"].pop()
        (self.directory / "source_manifest.json").write_text(json.dumps(self.manifest), encoding="utf-8")
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.verify_source_hashes(self.directory)

    def test_municipality_and_department_codes_must_correspond(self):
        path = self.directory / "municipalities.html"
        path.write_text(path.read_text(encoding="utf-8").replace("<td>05</td>", "<td>08</td>"), encoding="utf-8")
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.load_sources(self.directory)

    def test_unknown_calendar_boolean_is_rejected(self):
        path = self.directory / "calendar_2025.xlsx"
        frame = pd.read_excel(path)
        frame["is_holiday"] = frame["is_holiday"].astype(object)
        frame.loc[0, "is_holiday"] = "desconocido"
        frame.to_excel(path, sheet_name="calendar", index=False)
        with self.assertRaises(pipeline.DataQualityError):
            pipeline.load_sources(self.directory)


class SamplingTests(unittest.TestCase):
    def test_60_strata_have_20_reproducible_observations_each(self):
        rows = [
            {"city": city, "year": 2025, "month": month, "source_row": index}
            for index, (city, month, item) in enumerate(
                (city, month, item)
                for city in ["Barranquilla", "Bogota", "Bucaramanga", "Cali", "Medellin"]
                for month in range(1, 13) for item in range(25)
            )
        ]
        population = pd.DataFrame(rows)
        sample, coverage = pipeline.representative_sample(population)
        repeated, repeated_coverage = pipeline.representative_sample(population)
        other_seed, _ = pipeline.representative_sample(population, seed=7)
        self.assertEqual(len(sample), 1200)
        self.assertEqual(len(coverage), 60)
        self.assertTrue(sample.groupby(["city", "year", "month"]).size().eq(20).all())
        self.assertTrue(coverage["sample_weight"].eq(1.25).all())
        self.assertTrue(sample["source_row"].is_unique)
        pd.testing.assert_frame_equal(sample, repeated)
        pd.testing.assert_frame_equal(coverage, repeated_coverage)
        self.assertNotEqual(sample["source_row"].tolist(), other_seed["source_row"].tolist())

    def test_small_group_is_kept_in_full_with_weight_one(self):
        population = pd.DataFrame({
            "city": ["A"] * 3 + ["B"] * 25,
            "year": [2025] * 28, "month": [1] * 28, "source_row": range(28),
        })
        sample, coverage = pipeline.representative_sample(population)
        self.assertEqual(sample.loc[sample["city"].eq("A"), "source_row"].tolist(), [0, 1, 2])
        self.assertEqual(len(sample), 23)
        small = coverage.loc[coverage["city"].eq("A")].iloc[0]
        self.assertEqual((small["population_rows"], small["sample_rows"], small["sample_weight"]), (3, 3, 1.0))

    def test_nonpositive_quota_is_rejected(self):
        frame = pd.DataFrame({"city": ["A"], "year": [2025], "month": [1], "source_row": [1]})
        for quota in [0, -1]:
            with self.subTest(quota=quota), self.assertRaises(ValueError):
                pipeline.representative_sample(frame, quota=quota)


class EndToEndTests(unittest.TestCase):
    def test_small_run_preserves_database_and_exports_equivalent_csv_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            sources = directory / "sources"
            manifest = make_sources(sources)
            database = directory / "fixture.db"
            write_database(example_base(), database)
            original_hash = pipeline.sha256(database)
            output = directory / "output"
            with redirect_stdout(StringIO()):
                summary = pipeline.run(database, sources, output)
            self.assertEqual(pipeline.sha256(database), original_hash)
            self.assertEqual((summary["rows_before"], summary["rows_after"], summary["sample_rows"]), (3, 3, 3))
            self.assertEqual(summary["columns_before"], 32)
            self.assertEqual(summary["incomplete_rows"], 0)
            self.assertTrue(summary["original_values_unchanged"])
            exported = assert_exports_agree(self, output)
            self.assertTrue(exported["municipality_code"].eq("05001").all())
            self.assertTrue(exported["department_code"].eq("05").all())
            audit = output / "static/auditoria"
            self.assertEqual(json.loads((audit / "source_manifest.json").read_text(encoding="utf-8")), manifest)
            stored_summary = json.loads((audit / "integration_summary.json").read_text(encoding="utf-8"))
            self.assertEqual(stored_summary, summary)
            self.assertGreater((audit / "enriched_report.txt").stat().st_size, 0)
            for relative, expected_hash in summary["outputs_sha256"].items():
                self.assertEqual(pipeline.sha256(output / relative), expected_hash)

    @unittest.skipUnless(
        (PROJECT / "src/db/cleaned.duckdb").is_file()
        and all((PROJECT / "src/sources" / name).is_file() for name, _, _ in pipeline.SOURCE_CONTRACTS.values())
        and (PROJECT / "src/sources/source_manifest.json").is_file(),
        "La integración real requiere la base y las seis instantáneas locales del proyecto.",
    )
    def test_real_historical_sources_preserve_all_rows_and_sample_contract(self):
        database = PROJECT / "src/db/cleaned.duckdb"
        source_dir = PROJECT / "src/sources"
        input_hash = pipeline.sha256(database)
        before_hashes = {path.name: pipeline.sha256(path) for path in source_dir.iterdir() if path.is_file()}
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "output"
            with redirect_stdout(StringIO()):
                summary = pipeline.run(database, source_dir, output)
            self.assertEqual((summary["rows_before"], summary["rows_after"]), (43800, 43800))
            self.assertEqual((summary["columns_before"], summary["sample_rows"], summary["sample_strata"]), (32, 1200, 60))
            self.assertEqual(summary["incomplete_rows"], 0)
            self.assertEqual(summary["duplicate_city_hour_after"], 0)
            self.assertEqual(summary["synthetic_energy_rows"], 43800)
            full = pd.read_csv(output / "data/enriched_full.csv", keep_default_na=False, dtype={"municipality_code": str, "department_code": str})
            original = pipeline.load_base(database)
            expected = pipeline.serializable_frame(original.assign(local_date=pd.to_datetime(original["observation_time"].dt.date))).drop(columns="local_date")
            actual_original = pd.read_csv(StringIO(full[original.columns].to_csv(index=False)), keep_default_na=False)
            expected_original = pd.read_csv(StringIO(expected.to_csv(index=False)), keep_default_na=False)
            pd.testing.assert_frame_equal(actual_original, expected_original, check_exact=False, rtol=1e-12, atol=1e-10)
            self.assertTrue(full["year"].eq(2025).all())
            self.assertTrue(full.loc[full["city"].eq("Bucaramanga"), "divipola_year"].eq(2024).all())
            self.assertTrue(full.loc[full["city"].map(pipeline.normalize_city).eq("medellin"), "municipality_code"].eq("05001").all())
            coverage = pd.read_csv(output / "static/auditoria/sample_coverage.csv")
            self.assertTrue(coverage["sample_rows"].eq(20).all())
            self.assertEqual(int(coverage["population_rows"].sum()), 43800)
            assert_exports_agree(self, output)
            manifest = json.loads((output / "static/auditoria/source_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["period"], {"start": "2025-01-01", "end": "2025-12-31", "timezone": pipeline.TIMEZONE})
            self.assertFalse(manifest["network_required_for_pipeline"])
            self.assertEqual(len(manifest["sources"]), 6)
        self.assertEqual(pipeline.sha256(database), input_hash)
        self.assertEqual({path.name: pipeline.sha256(path) for path in source_dir.iterdir() if path.is_file()}, before_hashes)


if __name__ == "__main__":
    unittest.main()
